"""Document text extraction helpers for agent-readable previews."""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from html import unescape
from pathlib import Path
from typing import Any

DOCUMENT_TEXT_MAX_CHARACTERS = 30000
SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".csv",
    ".docx",
    ".html",
    ".htm",
    ".json",
    ".md",
    ".odt",
    ".pdf",
    ".pptx",
    ".rtf",
    ".txt",
    ".xlsx",
    ".xml",
}
TEMPORAL_MEDIA_PREFIXES = ("audio/", "video/")


@dataclass(frozen=True)
class DocumentExtractionResult:
    readable: bool
    text: str = ""
    truncated: bool = False
    page_or_sheet_count: int | None = None
    warnings: list[str] = field(default_factory=list)
    message: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "readable": self.readable,
            "text": self.text,
            "truncated": self.truncated,
            "page_or_sheet_count": self.page_or_sheet_count,
            "warnings": self.warnings,
            "message": self.message,
        }


def is_temporal_media_type(media_type: str | None) -> bool:
    normalized = str(media_type or "").lower()
    return normalized.startswith(TEMPORAL_MEDIA_PREFIXES)


def is_supported_document_path(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS


def extract_document_text(
    file_path: str | Path,
    *,
    media_type: str | None = None,
    max_characters: int = DOCUMENT_TEXT_MAX_CHARACTERS,
) -> DocumentExtractionResult:
    path = Path(file_path)
    suffix = path.suffix.lower()
    warnings: list[str] = []

    if is_temporal_media_type(media_type):
        return DocumentExtractionResult(
            readable=False,
            message="Audio and video files are not supported. Do not infer their contents.",
        )

    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        return DocumentExtractionResult(
            readable=False,
            message="This file type is not directly readable. Do not infer its contents.",
        )

    try:
        text, count = _extract_by_extension(path, suffix)
    except Exception as exc:
        return DocumentExtractionResult(
            readable=False,
            warnings=[f"{type(exc).__name__}: {exc}"],
            message="Text extraction failed. Do not infer this document's contents.",
        )

    normalized = _normalize_text(text)
    if not normalized:
        message = (
            "No extractable text was found. This may be a scanned or image-only document; "
            "do not infer its contents."
        )
        return DocumentExtractionResult(
            readable=False,
            page_or_sheet_count=count,
            warnings=warnings,
            message=message,
        )

    truncated = len(normalized) > max_characters
    if truncated:
        normalized = normalized[:max_characters].rstrip()
        normalized = (
            "[Document text truncated to the first "
            f"{max_characters} characters.]\n\n{normalized}\n\n[End of truncated document text.]"
        )
        warnings.append(f"Document text truncated to {max_characters} characters.")

    return DocumentExtractionResult(
        readable=True,
        text=normalized,
        truncated=truncated,
        page_or_sheet_count=count,
        warnings=warnings,
        message="Document text extracted.",
    )


def _extract_by_extension(path: Path, suffix: str) -> tuple[str, int | None]:
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    if suffix == ".csv":
        return _extract_csv(path)
    if suffix in {".html", ".htm"}:
        return _extract_html(path)
    if suffix == ".json":
        return _extract_json(path)
    if suffix == ".rtf":
        return _extract_rtf(path)
    if suffix == ".odt":
        return _extract_odt(path)
    return path.read_text(encoding="utf-8", errors="replace"), None


def _extract_pdf(path: Path) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {index}]\n{text}")
    return "\n\n".join(pages), len(reader.pages)


def _extract_docx(path: Path) -> tuple[str, int | None]:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append("\t".join(values))
    return "\n".join(parts), None


def _extract_pptx(path: Path) -> tuple[str, int]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"[Slide {slide_index}]\n" + "\n".join(texts))
    return "\n\n".join(slides), len(presentation.slides)


def _extract_xlsx(path: Path) -> tuple[str, int]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append("\t".join(values).rstrip())
        if rows:
            sheets.append(f"[Sheet: {worksheet.title}]\n" + "\n".join(rows))
    return "\n\n".join(sheets), len(workbook.worksheets)


def _extract_csv(path: Path) -> tuple[str, None]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        rows = ["\t".join(row) for row in csv.reader(handle)]
    return "\n".join(rows), None


def _extract_html(path: Path) -> tuple[str, None]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text("\n"), None


def _extract_json(path: Path) -> tuple[str, None]:
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    return json.dumps(data, ensure_ascii=False, indent=2), None


def _extract_rtf(path: Path) -> tuple[str, None]:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(path.read_text(encoding="utf-8", errors="replace")), None


def _extract_odt(path: Path) -> tuple[str, None]:
    from odf import teletype
    from odf.opendocument import load
    from odf.text import P

    document = load(str(path))
    paragraphs = [teletype.extractText(node) for node in document.getElementsByType(P)]
    return "\n".join(paragraphs), None


def _normalize_text(text: str) -> str:
    lines = [unescape(line).strip() for line in str(text or "").replace("\r\n", "\n").split("\n")]
    collapsed: list[str] = []
    previous_blank = False
    for line in lines:
        blank = not line
        if blank and previous_blank:
            continue
        collapsed.append(line)
        previous_blank = blank
    return "\n".join(collapsed).strip()
