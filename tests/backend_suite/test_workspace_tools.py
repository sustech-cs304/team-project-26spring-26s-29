"""Tests for workspace file tools and command helpers."""

import json
from pathlib import Path

from backend.agent.context import WorkspaceInfoProvider
from backend.agent.instructions import build_agent_instructions
from backend.agent.tools import (
    WORKSPACE_TOOLS,
    create_workspace_file,
    list_workspace_files,
    preview_workspace_file_tool_impl,
    read_workspace_file_tool_impl,
    run_workspace_python_tool_impl,
    run_workspace_shell_tool_impl,
    search_workspace_text_tool_impl,
    send_workspace_file_tool_impl,
    update_workspace_file,
)
from backend.services import resolve_workspace_path
from backend.services.document_service import extract_document_text

from .support import AsyncBackendTestCase, BackendTestCase


class WorkspaceToolTests(BackendTestCase):
    def test_workspace_file_tools_round_trip(self) -> None:
        created = create_workspace_file("outputs/report.txt", "hello")
        updated = update_workspace_file("outputs/report.txt", "\nworld", mode="append")
        listed = list_workspace_files("outputs")
        read_back = read_workspace_file_tool_impl("outputs/report.txt")
        searched = search_workspace_text_tool_impl("world", "outputs")

        self.assertEqual(created["action"], "create")
        self.assertEqual(updated["action"], "append")
        self.assertEqual(read_back["text"], "hello\nworld")
        self.assertEqual(searched["matches"][0]["relative_path"], "outputs/report.txt")
        self.assertEqual(listed["entries"][0]["relative_path"], "outputs/report.txt")

    def test_workspace_tools_reject_path_escape(self) -> None:
        with self.assertRaises(ValueError):
            read_workspace_file_tool_impl("../outside.txt")

    def test_workspace_command_tools_run_in_workspace(self) -> None:
        shell_result = run_workspace_shell_tool_impl("Write-Output hello")
        python_result = run_workspace_python_tool_impl("print('workspace-python')")

        self.assertEqual(shell_result["exit_code"], 0)
        self.assertIn("hello", shell_result["stdout"])
        self.assertEqual(python_result["exit_code"], 0)
        self.assertIn("workspace-python", python_result["stdout"])

    def test_preview_workspace_file_returns_rich_preview_items(self) -> None:
        create_workspace_file("outputs/preview.txt", "hello preview world")
        pdf_path = resolve_workspace_path("outputs/sample.pdf")
        pdf_path.write_bytes(_sample_pdf_bytes("hello pdf preview"))

        text_preview = preview_workspace_file_tool_impl("outputs/preview.txt")
        pdf_preview = preview_workspace_file_tool_impl("outputs/sample.pdf")

        self.assertEqual(text_preview[0].type, "text")
        self.assertEqual(text_preview[1].type, "text")
        self.assertIn("hello preview world", text_preview[1].text)
        self.assertEqual(pdf_preview[1].type, "text")
        self.assertIn("hello pdf preview", pdf_preview[1].text)

    def test_send_workspace_file_returns_downloadable_file_item(self) -> None:
        pptx_path = resolve_workspace_path("outputs/slides.pptx")
        _write_pptx(pptx_path, "hello sendable pptx")

        result = send_workspace_file_tool_impl("outputs/slides.pptx")

        self.assertEqual(result[0].type, "text")
        self.assertIn("Workspace file ready to send", result[0].text)
        self.assertEqual(result[1].type, "uri")
        self.assertEqual(result[1].media_type, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        self.assertEqual(result[1].additional_properties["name"], "slides.pptx")
        self.assertEqual(result[1].additional_properties["relativePath"], "outputs/slides.pptx")
        self.assertEqual(result[1].additional_properties["sizeBytes"], pptx_path.stat().st_size)

    def test_send_workspace_file_rejects_images(self) -> None:
        image_path = resolve_workspace_path("outputs/image.png")
        image_path.write_bytes(
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
            b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
            b"\x1f\x15\xc4\x89"
        )

        with self.assertRaisesRegex(ValueError, "preview_workspace_file"):
            send_workspace_file_tool_impl("outputs/image.png")

    def test_document_extraction_supports_common_formats_and_truncation(self) -> None:
        _write_docx(resolve_workspace_path("outputs/sample.docx"), "hello docx")
        _write_pptx(resolve_workspace_path("outputs/sample.pptx"), "hello pptx")
        _write_xlsx(resolve_workspace_path("outputs/sample.xlsx"), "hello xlsx")
        create_workspace_file("outputs/sample.csv", "name,value\nhello,csv")
        create_workspace_file("outputs/sample.html", "<html><body><h1>hello html</h1></body></html>")
        create_workspace_file("outputs/sample.json", json.dumps({"hello": "json"}))
        create_workspace_file("outputs/sample.txt", "hello txt")

        cases = {
            "outputs/sample.docx": "hello docx",
            "outputs/sample.pptx": "hello pptx",
            "outputs/sample.xlsx": "hello xlsx",
            "outputs/sample.csv": "hello\tcsv",
            "outputs/sample.html": "hello html",
            "outputs/sample.json": '"hello": "json"',
            "outputs/sample.txt": "hello txt",
        }
        for relative_path, expected in cases.items():
            extraction = extract_document_text(resolve_workspace_path(relative_path))
            self.assertTrue(extraction.readable, relative_path)
            self.assertIn(expected, extraction.text)

        create_workspace_file("outputs/long.txt", "x" * 31000)
        long_result = extract_document_text(resolve_workspace_path("outputs/long.txt"))
        self.assertTrue(long_result.readable)
        self.assertTrue(long_result.truncated)
        self.assertIn("Document text truncated", long_result.text)

    def test_empty_pdf_returns_unreadable_warning(self) -> None:
        from pypdf import PdfWriter

        pdf_path = resolve_workspace_path("outputs/empty.pdf")
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with pdf_path.open("wb") as handle:
            writer.write(handle)

        extraction = extract_document_text(pdf_path)

        self.assertFalse(extraction.readable)
        self.assertIn("No extractable text", extraction.message)

    def test_read_workspace_file_extracts_documents_and_warns_on_binary(self) -> None:
        pdf_path = resolve_workspace_path("outputs/readable.pdf")
        pdf_path.write_bytes(_sample_pdf_bytes("workspace pdf text"))
        binary_path = resolve_workspace_path("outputs/archive.bin")
        binary_path.write_bytes(bytes([0, 1, 2]))

        pdf_result = read_workspace_file_tool_impl("outputs/readable.pdf")
        binary_result = read_workspace_file_tool_impl("outputs/archive.bin")

        self.assertTrue(pdf_result["is_document"])
        self.assertIn("workspace pdf text", pdf_result["text"])
        self.assertFalse(binary_result["is_text"])
        self.assertIn("not directly readable", binary_result["message"])

    def test_workspace_tool_approval_modes(self) -> None:
        tool_modes = {tool.name: tool.approval_mode for tool in WORKSPACE_TOOLS}
        self.assertEqual(tool_modes["list_workspace_files"], "never_require")
        self.assertEqual(tool_modes["search_workspace_text"], "never_require")
        self.assertEqual(tool_modes["read_workspace_file"], "never_require")
        self.assertEqual(tool_modes["preview_workspace_file"], "never_require")
        self.assertEqual(tool_modes["send_workspace_file"], "never_require")
        self.assertEqual(tool_modes["create_workspace_file"], "always_require")
        self.assertEqual(tool_modes["update_workspace_file"], "always_require")
        self.assertEqual(tool_modes["run_workspace_shell"], "always_require")
        self.assertEqual(tool_modes["run_workspace_python"], "always_require")


class WorkspaceContextTests(AsyncBackendTestCase):
    async def test_workspace_provider_injects_snapshot(self) -> None:
        provider = WorkspaceInfoProvider()

        class DummyContext:
            def __init__(self) -> None:
                self.metadata = {}
                self.instructions = []

            def extend_instructions(self, source_id, text) -> None:
                self.instructions.append((source_id, text))

        context = DummyContext()
        session = type("Session", (), {"session_id": "ws-1", "state": {}})()
        state = {}

        await provider.before_run(agent=object(), session=session, context=context, state=state)

        self.assertIn("workspace_snapshot", state)
        self.assertIn("workspace_info", context.metadata)
        self.assertIn("Workspace context:", context.instructions[0][1])
        self.assertIn("Markdown image URLs", context.instructions[0][1])
        self.assertIn("Do not share non-image workspace files as Markdown links", context.instructions[0][1])
        self.assertIn("Do not delete uploaded inputs or generated outputs", context.instructions[0][1])

    async def test_agent_instructions_prefer_markdown_for_workspace_image_display(self) -> None:
        instructions = build_agent_instructions("en")

        self.assertIn("![description](outputs/image.png)", instructions)
        self.assertIn("Do not call preview_workspace_file only to display a workspace image", instructions)
        self.assertIn("use send_workspace_file for non-image files", instructions)
        self.assertIn("asks to download an image file, call preview_workspace_file", instructions)
        self.assertIn("Do not use preview_workspace_file to hand over non-image files", instructions)
        self.assertIn("Do not delete uploaded inputs or generated workspace artifacts", instructions)
        self.assertIn("Delete workspace files only when the user explicitly asks", instructions)


def _sample_pdf_bytes(text: str) -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        (
            b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n"
        ),
        b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
        f"5 0 obj << /Length {len(content)} >> stream\n".encode("ascii")
        + content
        + b"\nendstream endobj\n",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(payload))
        payload.extend(obj)
    xref_at = len(payload)
    payload.extend(b"xref\n0 6\n")
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(b"trailer << /Root 1 0 R /Size 6 >>\n")
    payload.extend(f"startxref\n{xref_at}\n%%EOF\n".encode("ascii"))
    return bytes(payload)


def _write_docx(path: Path, text: str) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    document.save(str(path))


def _write_pptx(path: Path, text: str) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = text
    presentation.save(str(path))


def _write_xlsx(path: Path, text: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = text
    workbook.save(str(path))
